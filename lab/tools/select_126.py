"""把「试验记录」那 126 条对应到 server/logs 的原始日志，产出专属索引。

用法:
    python tools/select_126.py
    python tools/select_126.py --xlsx ../../试验记录_20260923-20260924_含最终质量.xlsx \
                               --logs ../../server/logs --out data/runs_126.csv

产出:
    data/runs_126.csv   每条一行：xlsx 标签 + 遥测文件路径 + 遥测可用性

目的:
    后续所有分析/回放只在这 126 条上做，不碰全量 2630 个 run。
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys

LAB = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_XLSX = os.path.join(os.path.dirname(LAB), "..", "试验记录_20260923-20260924_含最终质量.xlsx")
DEFAULT_LOGS = os.path.join(os.path.dirname(LAB), "server", "logs")


def read_xlsx(path: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb["试验记录"]
    out = []
    for r in range(3, ws.max_row + 1):
        v = [ws.cell(r, c).value for c in range(1, 16)]
        if all(x is None for x in v):
            continue
        out.append(dict(
            seq=v[0], run_id=v[1],
            date=str(v[2])[:10] if v[2] else "",
            time=str(v[3]) if v[3] else "",
            powder=v[4], kind=v[5], result=v[6],
            target_mg=v[7], final_mg=v[8], dev_mg=v[9], dev_rate=v[10],
            duration_s=v[11], note=v[12],
        ))
    return out


def find_triple(logs: str, run_id: str) -> dict:
    """定位一个 run 的三件套：配置 JSON / 逐拍遥测 / 停后沉降。"""
    files = os.listdir(logs)
    js = [f for f in files if f.endswith(".json") and run_id in f]
    smp = [f for f in files if run_id in f and f.endswith("_samples.csv")
           and "_post_stop_samples.csv" not in f]
    post = [f for f in files if run_id in f and f.endswith("_post_stop_samples.csv")]
    if not js:
        return {}
    jp = os.path.join(logs, js[0])
    try:
        j = json.load(open(jp, encoding="utf-8-sig"))
    except Exception:
        return {"json_file": js[0]}
    samples = j.get("samples") or []
    keys = set()
    for s in samples[:1]:
        keys |= set(s.keys())
    for s in samples[-1:]:
        keys |= set(s.keys())
    return dict(
        json_file=js[0],
        samples_file=smp[0] if smp else "",
        post_stop_file=post[0] if post else "",
        n_samples=len(samples),
        algorithm_version=j.get("algorithm_version"),
        powder_id=j.get("powder_id"),
        stop_reason=(j.get("stop_reason") or "")[:60],
        has_rate=("control_rate_mg_s" in keys),
        has_proj=("projected_stop_mass_mg" in keys),
        has_tail=("estimated_tail_mg" in keys),
        has_confirm=("stop_confirmation_required" in keys),
        pred_cfg=json.dumps(j.get("predictive_stop_config") or {}, ensure_ascii=False),
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--logs", default=DEFAULT_LOGS)
    ap.add_argument("--out", default=os.path.join(LAB, "data", "runs_126.csv"))
    args = ap.parse_args()

    rows = read_xlsx(args.xlsx)
    print("xlsx 明细条数:", len(rows))

    miss = []
    enriched = []
    for r in rows:
        t = find_triple(args.logs, r["run_id"])
        if not t:
            miss.append(r["run_id"])
        merged = dict(r)
        merged.update(t)
        merged["samples_path"] = os.path.join("..", "..", "server", "logs", t["samples_file"]) if t.get("samples_file") else ""
        enriched.append(merged)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    fields = ["seq", "run_id", "date", "time", "powder", "powder_id", "kind", "result",
              "target_mg", "final_mg", "dev_mg", "dev_rate", "duration_s",
              "n_samples", "algorithm_version", "has_rate", "has_proj", "has_tail",
              "has_confirm", "json_file", "samples_file", "samples_path",
              "stop_reason", "note", "pred_cfg"]
    with open(args.out, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(enriched)

    have_smp = sum(1 for r in enriched if r.get("samples_file"))
    have_rate = sum(1 for r in enriched if r.get("has_rate"))
    have_proj = sum(1 for r in enriched if r.get("has_proj"))
    have_conf = sum(1 for r in enriched if r.get("has_confirm"))
    tot_samples = sum(r.get("n_samples") or 0 for r in enriched)
    print("写入:", args.out)
    print("有配置 JSON : %d / %d" % (len(enriched) - len(miss), len(enriched)))
    print("有逐拍遥测  : %d / %d" % (have_smp, len(enriched)))
    print("遥测含控制速率 control_rate_mg_s : %d" % have_rate)
    print("遥测含预测落点 projected_stop_mass_mg : %d" % have_proj)
    print("遥测含实际确认数 stop_confirmation_required : %d" % have_conf)
    print("遥测总拍数  : %d（平均 %.1f 拍/条）" % (tot_samples, tot_samples / max(1, len(enriched))))
    import collections
    print("算法版本分布:", dict(collections.Counter(r.get("algorithm_version") for r in enriched)))
    if miss:
        print("\n未找到日志的 run（%d 条）:" % len(miss))
        for m in miss[:10]:
            print("   ", m)
    return 0


if __name__ == "__main__":
    sys.exit(main())
