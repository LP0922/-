"""Build the phase-one powder dispensing technical and collaboration proposal."""

from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "output" / "阶段一粉末加样软件技术与协作方案.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
RISK_RED = "9B1C1C"
TEXT = "1F2937"


def set_font(run, name: str = "Microsoft YaHei", size: float = 11, bold: bool = False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int]):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for idx, width in enumerate(widths_dxa):
        grid.gridCol_lst[idx].set(qn("w:w"), str(width))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            margins = tc_pr.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for side, value in (("top", "80"), ("bottom", "80"), ("start", "120"), ("end", "120")):
                node = margins.find(qn(f"w:{side}"))
                if node is None:
                    node = OxmlElement(f"w:{side}")
                    margins.append(node)
                node.set(qn("w:w"), value)
                node.set(qn("w:type"), "dxa")


def format_table(table, widths: list[int], header=True):
    set_table_geometry(table, widths)
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_font(run, size=9.5, bold=header and row_index == 0, color=TEXT)
            if header and row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
    if header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        header_element = OxmlElement("w:tblHeader")
        header_element.set(qn("w:val"), "true")
        tr_pr.append(header_element)


def add_heading(document, text: str, level: int):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(text)
    return paragraph


def add_body(document, text: str, bold_lead: str | None = None):
    paragraph = document.add_paragraph()
    if bold_lead:
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True, color=DARK_BLUE)
    run = paragraph.add_run(text)
    set_font(run, color=TEXT)
    return paragraph


def add_bullets(document, items: list[str]):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_font(run, color=TEXT)


def add_numbered(document, items: list[str]):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_font(run, color=TEXT)


def add_callout(document, label: str, text: str, color: str = LIGHT_GRAY):
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    paragraph = cell.paragraphs[0]
    label_run = paragraph.add_run(label)
    set_font(label_run, size=10.5, bold=True, color=DARK_BLUE if color != "FDECEC" else RISK_RED)
    text_run = paragraph.add_run(text)
    set_font(text_run, size=10.5, color=TEXT)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(document, headers: list[str], rows: list[list[str]], widths: list[int]):
    table = document.add_table(rows=1, cols=len(headers))
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            cells[idx].text = value
    format_table(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_styles(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for level, size, color, before, after in (
        (1, 16, BLUE, 16, 8),
        (2, 13, BLUE, 12, 6),
        (3, 12, DARK_BLUE, 8, 4),
    ):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("阶段一粉末加样软件技术与协作方案")
    set_font(header_run, size=8.5, color="6B7280")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("内部技术讨论稿 | " + str(date.today()))
    set_font(footer_run, size=8.5, color="6B7280")


def build_document() -> Document:
    document = Document()
    configure_styles(document)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(72)
    title.paragraph_format.space_after = Pt(14)
    title_run = title.add_run("阶段一粉末加样软件技术与协作方案")
    set_font(title_run, size=24, bold=True, color=BLUE)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    subtitle_run = subtitle.add_run("面向 0.1 g 加粉、设备联调与数据驱动优化的初版方案")
    set_font(subtitle_run, size=12, color="6B7280")

    meta = document.add_table(rows=4, cols=2)
    for idx, (key, value) in enumerate(
        [
            ("阶段目标", "稳定完成 0.1 g（100 mg）加粉，并建立可追溯数据闭环"),
            ("当前设备", "AT8811C 称量模块；ESP32-S3 / LA10-D 电缸网关"),
            ("算法策略", "L0 基线 + 四段式控制 + 尾量估计 + ILC 影子模式"),
            ("协作方式", "设备集成与在线执行 / 数据格式与离线算法并行"),
        ]
    ):
        meta.cell(idx, 0).text = key
        meta.cell(idx, 1).text = value
    format_table(meta, [2400, 6960], header=False)
    for row in meta.rows:
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for run in row.cells[0].paragraphs[0].runs:
            set_font(run, size=10, bold=True, color=DARK_BLUE)

    document.add_page_break()

    add_heading(document, "1. 项目目标与阶段边界", 1)
    add_body(document, "阶段一的主目标是建立 0.1 g（100 mg）粉末加样的基础软件闭环，使每次任务具备可控制、可验收、可追溯和可复盘能力。支线目标是在不降低安全性的前提下，识别当前机构、粉体和称量系统可达到的实际精度极限。")
    add_body(document, "阶段一不以 GPR、贝叶斯优化等高级模型为上线目标。系统先使用可解释的 L0 参数表、状态机和硬性安全边界完成稳定运行，再以真实过程数据支持后续模型优化。")
    add_callout(document, "阶段一原则：", "实时安全、互锁、超量保护、通信异常和人工急停优先于任何算法输出。ILC 只影响后续任务的候选前馈参数，不能绕过在线安全链路。")

    add_heading(document, "2. 当前设备基础与约束", 1)
    add_table(
        document,
        ["设备", "已具备能力", "阶段一用途", "当前约束"],
        [
            ["AT8811C", "Modbus RTU 重量、稳定/零区状态、受授权清零", "称量、稳定判定、最终验收与过程记录", "已确认寄存器读数单位为 mg；当前实测存在明显零点漂移，需作为质量门限处理"],
            ["ESP32-S3 / LA10-D", "使能、相对位移、暂停、急停、故障与状态读取", "后续作为给粉机构中的执行器或门控动作来源", "尚需确认 LA10 对应的实际机械动作与出粉关系"],
        ],
        [1500, 2500, 2500, 2860],
    )
    add_body(document, "当前驱动解决的是设备通信与命令确认，不等同于完整给粉能力。只有确认 LA10 或后续振动器、门控、螺杆等设备的动作语义，并完成“动作量到质量增量”的标定后，才能形成真实 Feeder 执行器。")
    add_callout(document, "当前风险：", "天平空载和带载测试均观察到 mg 级到数十 mg 级变化。任务可以继续采集数据，但在稳定窗口、基线漂移和计量标定未冻结前，不得将这些数据用于自动 ILC 更新。", "FDECEC")

    add_heading(document, "3. 总体软件架构", 1)
    diagram = document.add_paragraph()
    diagram.paragraph_format.space_after = Pt(8)
    run = diagram.add_run(
        "配方 / CandidatePlan\n"
        "        |\n"
        "在线安全状态机 -> 四段式控制器 -> Feeder 执行器 -> 真实设备\n"
        "        |                   |                |\n"
        "        |                   |                +-> LA10 / 后续门控、振动、互锁\n"
        "        |                   +-> L0 参数、尾量估计、稳定判定\n"
        "        +-> RunRecord <----- AT8811C 称量时间序列\n"
        "                              |\n"
        "                    离线回放 / ILC 影子分析 / 参数优化"
    )
    set_font(run, name="Consolas", size=9.5, color=DARK_BLUE)

    add_table(
        document,
        ["层级", "核心模块", "职责"],
        [
            ["设备适配", "AT8811C、LA10、Feeder、互锁", "通信、命令确认、状态读取和硬件故障返回；不做工艺决策"],
            ["在线编排", "任务状态机、安全策略", "执行、暂停、重试、超时、急停和安全退出"],
            ["控制算法", "L0、四段控制、尾量估计、稳定判定", "决定切段、动作上限和何时停止给粉"],
            ["数据与学习", "RunRecord、回放、ILC 影子", "记录事实、验证样本质量、给出下一次候选参数"],
        ],
        [1500, 2800, 5060],
    )

    add_heading(document, "4. 四段式控制与状态机", 1)
    add_body(document, "单次任务由状态机控制，算法不直接下发串口帧。每一段都受最大动作量、最大累计补偿、最大等待时间和硬性超量边界约束。")
    add_table(
        document,
        ["阶段", "控制目标", "主要决策依据", "退出条件"],
        [
            ["PRECHECK / TARE", "确认设备可用并建立任务基线", "通信、互锁、称量质量、配方和清零授权", "通过后才可开始给粉"],
            ["粗加", "快速接近目标，预留安全余量", "L0 前馈动作量与当前质量", "达到粗加停料阈值或触发保护"],
            ["缓补", "降低速度，减少尾量风险", "预测剩余量、近期质量变化、动作上限", "进入精加余量区间"],
            ["精加", "小脉冲逐步逼近目标", "每脉冲后称量、稳定窗口、最大脉冲数", "合格、无法安全补偿或超时"],
            ["收尾验收", "等待余粉落料并确认最终质量", "硬件稳定标志 + 软件波动/漂移窗口", "ACCEPT、REJECT 或人工复核"],
        ],
        [1500, 2600, 3000, 2260],
    )

    add_heading(document, "5. L0、尾量估计与 ILC 的上线策略", 1)
    add_body(document, "L0 是阶段一唯一可投产的前馈基线。每条参数按粉体型号、粉体批次、给粉头型、目标质量区间和配方版本组织，包含粗加动作量、缓补上限、精加脉冲、收尾等待时间和硬性质量边界。")
    add_body(document, "停料后不应把瞬时秤值直接作为最终质量。控制器使用当前质量、余粉估计和短时漂移计算预测最终质量；预测只用于提前切段和停止给粉，最终验收始终以稳定实测质量为准。")
    add_table(
        document,
        ["层级", "阶段一状态", "允许做什么", "禁止做什么"],
        [
            ["L0", "上线基线", "查表、线性插值、人工审批参数", "越过动作/超量/超时安全边界"],
            ["尾量估计", "上线辅助", "辅助切段、决定最短静置时间", "替代最终实测验收"],
            ["ILC 影子", "数据积累", "离线生成建议、评估误差趋势", "直接修改当前任务或自动下发"],
            ["受限 ILC", "后续阶段", "经审批后限幅更新 L1", "使用异常、漂移超限或人工干预样本"],
        ],
        [1500, 1900, 3150, 2810],
    )
    add_callout(document, "学习样本准入：", "只有通信完整、稳定质量可信、无堵料/低流/超量/人工干预、配置可追溯且未触发天平质量门限的 RunRecord，才允许进入 ILC 数据集。")

    add_heading(document, "6. 数据合同与可追溯性", 1)
    add_body(document, "协作双方通过两个版本化对象连接。算法侧只提交 CandidatePlan，实机侧只返回事实性的 RunRecord。这样可避免算法直接操控串口，也能保证离线回放与实际执行使用同一套数据语义。")
    add_table(
        document,
        ["对象", "由谁提供", "关键内容", "用途"],
        [
            ["CandidatePlan", "数据与算法负责人", "目标 mg、阶段动作计划、参数版本、安全边界、最长时长", "提交在线执行器；不包含 COM 口和 Modbus 帧"],
            ["RunRecord", "设备集成负责人", "重量时间序列、状态切换、动作命令、报警、设备配置、最终结果", "数据校验、回放、L0/ILC 分析与审计"],
        ],
        [1700, 1900, 3650, 2110],
    )
    add_bullets(
        document,
        [
            "weight_trace 必须同时保存重量 mg、原始计数、稳定/零区位、时间戳和通信结果。",
            "command_trace 必须保存计划动作、实际下发值、开始/结束时间、确认结果与执行器状态。",
            "parameter_snapshot 必须保存配方、L0/L1、设备配置、算法版本和配置哈希。",
            "发生人工干预、通信异常或安全退出时，必须保留事件原因，不能只保留最终重量。",
        ],
    )

    add_heading(document, "7. 双人协作分工", 1)
    add_table(
        document,
        ["工作项", "A：设备集成与在线执行", "B：数据与算法", "共同确认"],
        [
            ["设备与安全", "驱动、Feeder 映射、互锁、急停、在线状态机", "定义算法所需观测与动作语义", "动作单位、设备边界和报警码"],
            ["标定与数据", "执行空载/带载/标准砝码/动作量实验，记录 RunRecord", "定义数据格式、校验质量、分析漂移和动作增益", "实验计划、样本是否可用"],
            ["控制参数", "加载经批准的参数版本，保证动作受限", "构建 L0、尾量估计和 ILC 影子建议", "误差目标、限幅与回退阈值"],
            ["实机联调", "占用 COM 口、执行测试、处理硬件故障", "离线回放、提出下一版 CandidatePlan", "每轮结论、版本发布或回退"],
        ],
        [1450, 2850, 2850, 2210],
    )
    add_body(document, "实机测试窗口由 A 负责。B 不与 A 同时占用 COM 口，而是使用已归档的 RunRecord 开展回放、参数分析和模拟。公共接口、配方结构和安全阈值变更必须双方评审。")

    add_heading(document, "8. 实施阶段与交付物", 1)
    add_table(
        document,
        ["阶段", "A 的交付", "B 的交付", "退出条件"],
        [
            ["0. 接口冻结", "设备清单、动作范围、互锁和通讯配置", "CandidatePlan / RunRecord v1、单位与时间规则", "双方签字确认数据合同"],
            ["1. 基础可观测", "驱动自检、采样、日志落盘", "漂移与稳定性分析、数据校验器", "形成有效 RunRecord 样例"],
            ["2. 原子动作标定", "Feeder 映射、单动作/停料实验", "动作量-质量、延迟和尾量模型", "L0 候选表和动作安全上限"],
            ["3. 固定 L0 闭环", "完整安全状态机和受控实机运行", "回放、稳定窗口、验收规则", "100 mg 端到端测试报告"],
            ["4. ILC 影子", "完整参数与过程归档", "样本准入、ILC 建议和效果报告", "决定是否批准受限 L1"],
        ],
        [1500, 2800, 2850, 2210],
    )

    add_heading(document, "9. 当前待冻结事项", 1)
    add_numbered(
        document,
        [
            "确认 LA10 在加粉机构中的具体角色，以及后续振动器、门控、传感器和互锁的接口。",
            "完成 AT8811C 的预热、清零、空载/带载、标准砝码与启动置零配置测试，冻结称量质量门限。",
            "确定 100 mg 配方的允许欠量、允许过量、任务最长时长、重试次数和失败处置。",
            "冻结 Feeder 动作单位的业务含义，例如步数、毫秒、开度或组合动作。",
            "建立首版 CandidatePlan、RunRecord 与 L0 参数表，先进入固定 L0 闭环实验。",
        ],
    )
    add_callout(document, "阶段一验收口径：", "应至少覆盖 100 mg 装样成功率、欠量率、过量及最大过量、最终误差和重复性、平均耗时、稳定判定耗时、异常识别率、RunRecord 完整率，以及 ILC 影子建议的离线收益。")

    add_heading(document, "10. 结论", 1)
    add_body(document, "本方案以 L0、四段状态机、称量质量门限和完整 RunRecord 为阶段一基础。设备侧先保证安全、可执行和可记录；算法侧先保证可回放、可解释和可限幅。当天平质量和给粉动作标定达到要求后，再由 ILC 影子数据逐步推动参数优化，而不是在设备能力未确认时直接引入复杂模型。")
    return document


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
